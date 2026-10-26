"""把講稿 Markdown（標題、段落、引言、清單、表格、粗體、行內程式碼、分隔線）轉成 Word 檔。
用法：python presentation/md2docx.py <輸入.md> <輸出.docx>（例：presentation/講稿與背景說明.md、presentation/評審問答演練.md）
"""
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Microsoft JhengHei"
GREEN = RGBColor(0x2C, 0x5F, 0x2D)
GRAY = RGBColor(0x5E, 0x6B, 0x63)


def set_font(run, size=None, bold=None, color=None, mono=False):
    f = "Consolas" if mono else FONT
    run.font.name = f
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(k), FONT if (k == "w:eastAsia" or not mono) else f)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_inline(p, text, size=11, color=None, italic=False):
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            r = p.add_run(part[2:-2])
            set_font(r, size, True, color)
        elif part.startswith("`"):
            r = p.add_run(part[1:-1])
            set_font(r, size - 1, None, color, mono=True)
        else:
            r = p.add_run(part)
            set_font(r, size, None, color)
        r.italic = italic


def shade(cell, hex_):
    tcpr = cell._element.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hex_)
    tcpr.append(sh)


def left_border(p, hex_="2C5F2D"):
    ppr = p._element.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:left")
    for k, v in (("w:val", "single"), ("w:sz", "18"), ("w:space", "8"), ("w:color", hex_)):
        b.set(qn(k), v)
    bdr.append(b)
    ppr.append(bdr)


def bottom_rule(p):
    ppr = p._element.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "C8D3CA")):
        b.set(qn(k), v)
    bdr.append(b)
    ppr.append(bdr)


def table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    head, body = cells[0], [r for r in cells[2:]]
    t = doc.add_table(rows=1 + len(body), cols=len(head))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate([head] + body):
        for j, txt in enumerate(row):
            c = t.cell(i, j)
            c.paragraphs[0].text = ""
            add_inline(c.paragraphs[0], txt, 10, RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else None)
            if i == 0:
                shade(c, "2C5F2D")
                for r in c.paragraphs[0].runs:
                    r.bold = True
    # 欄寬依各欄最長文字比例分配（總寬 16.6 cm）
    L = [max(len(r[j]) if j < len(r) else 0 for r in cells) for j in range(len(head))]
    L = [min(max(x, 4), 40) for x in L]
    t.autofit = False
    for j, x in enumerate(L):
        w = Cm(16.6 * x / sum(L))
        for row in t.rows:
            row.cells[j].width = w
    doc.add_paragraph()


def main(src, dst):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.25
    lines = open(src, encoding="utf-8").read().splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s:
            i += 1
            continue
        if s.startswith("|"):
            blk = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                blk.append(lines[i])
                i += 1
            table(doc, blk)
            continue
        if s == "---":
            bottom_rule(doc.add_paragraph())
        elif s.startswith("# "):
            p = doc.add_paragraph()
            add_inline(p, s[2:], 20, GREEN)
            for r in p.runs:
                r.bold = True
            p.paragraph_format.space_after = Pt(10)
        elif s.startswith("### "):
            p = doc.add_heading(level=2)
            p.text = ""
            add_inline(p, s[4:], 13, GREEN)
            for r in p.runs:
                r.bold = True
            p.paragraph_format.space_before = Pt(10)
        elif s.startswith("## "):
            p = doc.add_heading(level=1)
            p.text = ""
            add_inline(p, s[3:], 15, GREEN)
            for r in p.runs:
                r.bold = True
            p.paragraph_format.space_before = Pt(14)
        elif s.startswith("> "):
            p = doc.add_paragraph()
            left_border(p)
            p.paragraph_format.left_indent = Cm(0.4)
            add_inline(p, s[2:], 12)
        elif re.match(r"^\s*- ", ln):
            lvl = (len(ln) - len(ln.lstrip())) // 2
            p = doc.add_paragraph(style="List Bullet 2" if lvl else "List Bullet")
            add_inline(p, re.sub(r"^\s*- ", "", ln), 11)
        else:
            p = doc.add_paragraph()
            add_inline(p, s, 11, GRAY if s.startswith("簡報檔") else None)
        i += 1
    doc.core_properties.title = next((l[2:].strip() for l in lines if l.startswith("# ")), "決賽簡報")
    doc.core_properties.author = "TASTI_SatDashboard"
    doc.save(dst)
    print("saved", dst)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
